# src/data_staging/utils/file_handler.py - Complete File Handler Implementation

import pandas as pd
import json
import openpyxl
from pathlib import Path
from typing import Dict, Any, Optional, List, Union
import logging
import chardet
import csv
from datetime import datetime
import tempfile
import shutil

logger = logging.getLogger(__name__)

class FileHandler:
    """Utility class for handling various file formats"""
    
    def __init__(self):
        self.supported_formats = {
            '.csv': self._read_csv,
            '.xlsx': self._read_excel,
            '.xls': self._read_excel,
            '.json': self._read_json,
            '.parquet': self._read_parquet,
            '.txt': self._read_text
        }
    
    def read_file(self, file_path: Union[str, Path], chunksize: Optional[int] = None, **kwargs) -> Union[pd.DataFrame, Any]:
        """
        Read a file and return a pandas DataFrame or TextFileReader (iterator)
        
        Args:
            file_path: Path to the file
            chunksize: Number of rows per chunk (for CSV/Text). If provided, returns an iterator.
            **kwargs: Additional arguments for specific readers
        
        Returns:
            DataFrame, TextFileReader (iterator), or None if reading fails
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            logger.error(f"File not found: {file_path}")
            return None
        
        file_ext = file_path.suffix.lower()
        
        if file_ext not in self.supported_formats:
            logger.error(f"Unsupported file format: {file_ext}")
            return None
        
        try:
            logger.info(f"Reading file: {file_path} (chunksize={chunksize})")
            reader_func = self.supported_formats[file_ext]
            
            # Pass chunksize to reader function
            if chunksize:
                kwargs['chunksize'] = chunksize
                
            result = reader_func(file_path, **kwargs)
            
            if result is not None:
                if chunksize and file_ext in ['.csv', '.txt']:
                    logger.info(f"Successfully initiated chunked reading for {file_path}")
                elif isinstance(result, pd.DataFrame):
                    logger.info(f"Successfully read {len(result)} rows and {len(result.columns)} columns")
                    # If chunksize requested but format doesn't support streaming, yield full DF as one chunk
                    if chunksize:
                        return iter([result])
            
            return result
            
        except Exception as e:
            logger.error(f"Error reading file {file_path}: {e}")
            return None
    
    def _read_csv(self, file_path: Path, **kwargs) -> pd.DataFrame:
        """Read CSV file with intelligent encoding detection"""
        
        # Detect encoding
        encoding = self._detect_encoding(file_path)
        
        # Default CSV parameters
        csv_params = {
            'encoding': encoding,
            'sep': kwargs.get('delimiter', ','),
            'header': kwargs.get('header', 0),
            'skiprows': kwargs.get('skiprows', None),
            'nrows': kwargs.get('nrows', None),
            'usecols': kwargs.get('usecols', None),
            'dtype': kwargs.get('dtype', None),
            'parse_dates': kwargs.get('parse_dates', False),
            'date_parser': kwargs.get('date_parser', None),
            'na_values': kwargs.get('na_values', ['', 'NULL', 'null', 'N/A', 'n/a', 'NA']),
            'keep_default_na': True,
            'skipinitialspace': True,
            'chunksize': kwargs.get('chunksize', None),
            'on_bad_lines': kwargs.get('on_bad_lines', 'error')
        }
        
        # Try to detect delimiter if not specified
        if 'delimiter' not in kwargs:
            delimiter = self._detect_csv_delimiter(file_path, encoding)
            csv_params['sep'] = delimiter
        
        try:
            df = pd.read_csv(file_path, **csv_params)
            
            # If chunking, return iterator directly (cannot clean columns/rows here)
            if kwargs.get('chunksize'):
                return df
            
            # Clean column names
            df.columns = df.columns.str.strip()
            
            # Remove completely empty rows
            df = df.dropna(how='all')
            
            return df
            
        except UnicodeDecodeError:
            # Fallback to different encodings
            for fallback_encoding in ['latin-1', 'cp1252', 'iso-8859-1']:
                try:
                    csv_params['encoding'] = fallback_encoding
                    df = pd.read_csv(file_path, **csv_params)
                    logger.warning(f"Used fallback encoding {fallback_encoding} for {file_path}")
                    return df
                except:
                    continue
            raise
    
    def _read_excel(self, file_path: Path, **kwargs) -> pd.DataFrame:
        """Read Excel file"""
        
        excel_params = {
            'sheet_name': kwargs.get('sheet_name', 0),
            'header': kwargs.get('header', 0),
            'skiprows': kwargs.get('skiprows', None),
            'nrows': kwargs.get('nrows', None),
            'usecols': kwargs.get('usecols', None),
            'dtype': kwargs.get('dtype', None),
            'parse_dates': kwargs.get('parse_dates', False),
            'date_parser': kwargs.get('date_parser', None),
            'na_values': kwargs.get('na_values', ['', 'NULL', 'null', 'N/A', 'n/a', 'NA']),
            'keep_default_na': True
        }
        
        # Read Excel file
        if isinstance(excel_params['sheet_name'], str) or excel_params['sheet_name'] is None:
            df = pd.read_excel(file_path, **excel_params)
        else:
            # If sheet_name is int or list, handle accordingly
            df = pd.read_excel(file_path, **excel_params)
        
        # Clean column names
        if isinstance(df, pd.DataFrame):
            df.columns = df.columns.astype(str).str.strip()
            # Remove completely empty rows
            df = df.dropna(how='all')
        
        return df
    
    def _read_json(self, file_path: Path, **kwargs) -> pd.DataFrame:
        """Read JSON file"""
        
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Handle different JSON structures
        if isinstance(data, list):
            # Array of objects
            df = pd.json_normalize(data)
        elif isinstance(data, dict):
            # Check if it's a single record or nested structure
            if 'data' in data and isinstance(data['data'], list):
                # Common pattern: {"data": [...]}
                df = pd.json_normalize(data['data'])
            else:
                # Single object - convert to single-row DataFrame
                df = pd.json_normalize([data])
        else:
            raise ValueError(f"Unsupported JSON structure in {file_path}")
        
        if 'usecols' in kwargs and kwargs['usecols']:
            # Filter columns if requested
            available_cols = [c for c in kwargs['usecols'] if c in df.columns]
            if available_cols:
                df = df[available_cols]
        
        return df
    
    def _read_parquet(self, file_path: Path, **kwargs) -> pd.DataFrame:
        """Read Parquet file"""
        try:
            import pyarrow.parquet as pq
            # Map usecols to columns for read_parquet
            if 'usecols' in kwargs:
                kwargs['columns'] = kwargs.pop('usecols')
                
            df = pd.read_parquet(file_path, **kwargs)
            return df
        except ImportError:
            logger.error("PyArrow not installed. Cannot read Parquet files.")
            raise
    
    def _read_text(self, file_path: Path, **kwargs) -> pd.DataFrame:
        """Read text file (assume CSV-like format)"""
        # Default to tab-separated if no delimiter specified
        kwargs.setdefault('delimiter', '\t')
        return self._read_csv(file_path, **kwargs)
    
    def _detect_encoding(self, file_path: Path) -> str:
        """Detect file encoding"""
        try:
            with open(file_path, 'rb') as f:
                # Read first 10KB to detect encoding
                raw_data = f.read(10240)
                result = chardet.detect(raw_data)
                encoding = result['encoding']
                confidence = result['confidence']
                
                # If confidence is low, use utf-8 as fallback
                if confidence < 0.8:
                    logger.warning(f"Low encoding confidence ({confidence:.2f}) for {file_path}, using utf-8")
                    encoding = 'utf-8'
                
                logger.debug(f"Detected encoding: {encoding} (confidence: {confidence:.2f})")
                return encoding
                
        except Exception as e:
            logger.warning(f"Could not detect encoding for {file_path}: {e}")
            return 'utf-8'
    
    def _detect_csv_delimiter(self, file_path: Path, encoding: str) -> str:
        """Detect CSV delimiter"""
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                # Read first few lines
                sample = f.read(8192)
                
            # Use csv.Sniffer to detect delimiter
            sniffer = csv.Sniffer()
            delimiter = sniffer.sniff(sample).delimiter
            
            logger.debug(f"Detected delimiter: '{delimiter}'")
            return delimiter
            
        except Exception as e:
            logger.warning(f"Could not detect delimiter for {file_path}: {e}")
            return ','  # Default to comma
    
    def analyze_file_structure(self, file_path: Union[str, Path]) -> Dict[str, Any]:
        """Analyze file structure and return metadata"""
        file_path = Path(file_path)
        
        if not file_path.exists():
            return {"error": "File not found"}
        
        file_stats = file_path.stat()
        file_ext = file_path.suffix.lower()
        
        analysis = {
            "file_name": file_path.name,
            "file_size": file_stats.st_size,
            "file_extension": file_ext,
            "modified_time": datetime.fromtimestamp(file_stats.st_mtime).isoformat(),
            "supported": file_ext in self.supported_formats
        }
        
        if not analysis["supported"]:
            analysis["error"] = f"Unsupported file format: {file_ext}"
            return analysis
        
        try:
            # Read a sample of the file
            sample_df = self.read_file(file_path, nrows=100)
            
            if sample_df is not None:
                # Sanitize dataframe for JSON serialization (Replace NaN with None)
                # MUST cast to object first, otherwise None in float columns reverts to NaN
                sample_df_clean = sample_df.astype(object).where(pd.notnull(sample_df), None)
                
                analysis.update({
                    "columns": list(sample_df.columns),
                    "column_count": len(sample_df.columns),
                    "sample_row_count": len(sample_df),
                    "column_types": sample_df.dtypes.astype(str).to_dict(),
                    "sample_data": sample_df_clean.head(20).to_dict('records'),
                    "null_counts": sample_df.isnull().sum().to_dict(),
                    "estimated_total_rows": self._estimate_total_rows(file_path, file_ext)
                })
                
                # Add format-specific analysis
                if file_ext == '.xlsx':
                    analysis.update(self._analyze_excel_structure(file_path))
                elif file_ext == '.csv':
                    analysis.update(self._analyze_csv_structure(file_path))
                elif file_ext == '.json':
                    analysis.update(self._analyze_json_structure(file_path))
            else:
                analysis["error"] = "Could not read file"
                
        except Exception as e:
            analysis["error"] = f"Error analyzing file: {str(e)}"
        
        return analysis
    
    def _estimate_total_rows(self, file_path: Path, file_ext: str) -> int:
        """Estimate total number of rows in file"""
        try:
            if file_ext == '.csv':
                with open(file_path, 'r', encoding=self._detect_encoding(file_path)) as f:
                    # Count lines (rough estimate)
                    lines = sum(1 for _ in f)
                    return max(0, lines - 1)  # Subtract header
            
            elif file_ext in ['.xlsx', '.xls']:
                # For Excel, read without data to get shape
                df_info = pd.read_excel(file_path, nrows=0)
                workbook = openpyxl.load_workbook(file_path, read_only=True)
                sheet = workbook.active
                return sheet.max_row - 1 if sheet.max_row > 1 else 0
            
            elif file_ext == '.json':
                with open(file_path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    if isinstance(data, list):
                        return len(data)
                    elif isinstance(data, dict) and 'data' in data and isinstance(data['data'], list):
                        return len(data['data'])
                    else:
                        return 1
            
            return 0
            
        except Exception as e:
            logger.warning(f"Could not estimate row count for {file_path}: {e}")
            return 0
    
    def _analyze_excel_structure(self, file_path: Path) -> Dict[str, Any]:
        """Analyze Excel file structure"""
        try:
            excel_file = pd.ExcelFile(file_path)
            return {
                "excel_sheets": excel_file.sheet_names,
                "sheet_count": len(excel_file.sheet_names),
                "default_sheet": excel_file.sheet_names[0] if excel_file.sheet_names else None
            }
        except Exception as e:
            return {"excel_error": str(e)}
    
    def _analyze_csv_structure(self, file_path: Path) -> Dict[str, Any]:
        """Analyze CSV file structure"""
        try:
            encoding = self._detect_encoding(file_path)
            delimiter = self._detect_csv_delimiter(file_path, encoding)
            
            return {
                "csv_encoding": encoding,
                "csv_delimiter": delimiter,
                "csv_delimiter_name": self._get_delimiter_name(delimiter)
            }
        except Exception as e:
            return {"csv_error": str(e)}
    
    def _analyze_json_structure(self, file_path: Path) -> Dict[str, Any]:
        """Analyze JSON file structure"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            if isinstance(data, list):
                structure_type = "array"
                sample_keys = list(data[0].keys()) if data and isinstance(data[0], dict) else []
            elif isinstance(data, dict):
                structure_type = "object"
                sample_keys = list(data.keys())
            else:
                structure_type = "primitive"
                sample_keys = []
            
            return {
                "json_structure": structure_type,
                "json_keys": sample_keys[:10],  # First 10 keys
                "json_nested": self._has_nested_structure(data)
            }
        except Exception as e:
            return {"json_error": str(e)}
    
    def _get_delimiter_name(self, delimiter: str) -> str:
        """Get human-readable name for delimiter"""
        delimiter_names = {
            ',': 'comma',
            ';': 'semicolon',
            '\t': 'tab',
            '|': 'pipe',
            ' ': 'space'
        }
        return delimiter_names.get(delimiter, f"'{delimiter}'")
    
    def _has_nested_structure(self, data: Any, max_depth: int = 3, current_depth: int = 0) -> bool:
        """Check if JSON has nested structure"""
        if current_depth >= max_depth:
            return False
        
        if isinstance(data, dict):
            for value in data.values():
                if isinstance(value, (dict, list)):
                    return True
                if self._has_nested_structure(value, max_depth, current_depth + 1):
                    return True
        elif isinstance(data, list) and data:
            first_item = data[0]
            if isinstance(first_item, (dict, list)):
                return True
            if self._has_nested_structure(first_item, max_depth, current_depth + 1):
                return True
        
        return False
    
    def convert_file_format(
        self, 
        input_path: Union[str, Path], 
        output_path: Union[str, Path],
        output_format: str = 'csv',
        **kwargs
    ) -> bool:
        """
        Convert file from one format to another
        
        Args:
            input_path: Path to input file
            output_path: Path to output file
            output_format: Target format ('csv', 'excel', 'json', 'parquet')
            **kwargs: Additional arguments for readers/writers
        
        Returns:
            True if conversion successful, False otherwise
        """
        try:
            # Read input file
            df = self.read_file(input_path, **kwargs)
            
            if df is None:
                logger.error(f"Could not read input file: {input_path}")
                return False
            
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            
            # Write to output format
            if output_format.lower() == 'csv':
                df.to_csv(output_path, index=False, **kwargs)
            elif output_format.lower() == 'excel':
                df.to_excel(output_path, index=False, **kwargs)
            elif output_format.lower() == 'json':
                df.to_json(output_path, orient='records', **kwargs)
            elif output_format.lower() == 'parquet':
                df.to_parquet(output_path, index=False, **kwargs)
            else:
                logger.error(f"Unsupported output format: {output_format}")
                return False
            
            logger.info(f"Successfully converted {input_path} to {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"Error converting file: {e}")
            return False
    
    def create_backup(self, file_path: Union[str, Path], backup_dir: Optional[Union[str, Path]] = None) -> Optional[Path]:
        """
        Create a backup copy of a file
        
        Args:
            file_path: Path to file to backup
            backup_dir: Directory to store backup (default: same directory as original)
        
        Returns:
            Path to backup file or None if failed
        """
        try:
            file_path = Path(file_path)
            
            if not file_path.exists():
                logger.error(f"File not found: {file_path}")
                return None
            
            if backup_dir is None:
                backup_dir = file_path.parent
            else:
                backup_dir = Path(backup_dir)
                backup_dir.mkdir(parents=True, exist_ok=True)
            
            # Create backup filename with timestamp
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_name = f"{file_path.stem}_backup_{timestamp}{file_path.suffix}"
            backup_path = backup_dir / backup_name
            
            # Copy file
            shutil.copy2(file_path, backup_path)
            
            logger.info(f"Created backup: {backup_path}")
            return backup_path
            
        except Exception as e:
            logger.error(f"Error creating backup: {e}")
            return None
    
    def validate_file_integrity(self, file_path: Union[str, Path]) -> Dict[str, Any]:
        """
        Validate file integrity and readability
        
        Args:
            file_path: Path to file to validate
        
        Returns:
            Dictionary with validation results
        """
        file_path = Path(file_path)
        
        result = {
            "file_exists": file_path.exists(),
            "file_readable": False,
            "file_size": 0,
            "format_valid": False,
            "sample_readable": False,
            "estimated_rows": 0,
            "errors": []
        }
        
        if not result["file_exists"]:
            result["errors"].append("File does not exist")
            return result
        
        try:
            # Check file stats
            file_stats = file_path.stat()
            result["file_size"] = file_stats.st_size
            result["file_readable"] = True
            
            # Check if format is supported
            file_ext = file_path.suffix.lower()
            result["format_valid"] = file_ext in self.supported_formats
            
            if not result["format_valid"]:
                result["errors"].append(f"Unsupported file format: {file_ext}")
                return result
            
            # Try to read a sample
            sample_df = self.read_file(file_path, nrows=10)
            
            if sample_df is not None:
                result["sample_readable"] = True
                result["estimated_rows"] = self._estimate_total_rows(file_path, file_ext)
            else:
                result["errors"].append("Could not read file sample")
            
        except PermissionError:
            result["errors"].append("Permission denied")
        except Exception as e:
            result["errors"].append(f"Validation error: {str(e)}")
        
        return result

# Example usage and utility functions
def create_sample_files(output_dir: Union[str, Path]):
    """Create sample files for testing"""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # Sample data
    sample_data = {
        'id': [1, 2, 3, 4, 5],
        'name': ['Alice', 'Bob', 'Charlie', 'Diana', 'Eve'],
        'age': [25, 30, 35, 28, 32],
        'email': ['alice@example.com', 'bob@example.com', 'charlie@example.com', 'diana@example.com', 'eve@example.com'],
        'salary': [50000, 60000, 75000, 55000, 68000],
        'join_date': ['2023-01-15', '2022-06-20', '2021-03-10', '2023-08-05', '2022-11-30']
    }
    
    df = pd.DataFrame(sample_data)
    
    # Create different formats
    df.to_csv(output_dir / 'sample_data.csv', index=False)
    df.to_excel(output_dir / 'sample_data.xlsx', index=False)
    df.to_json(output_dir / 'sample_data.json', orient='records', indent=2)
    
    # Create a more complex JSON structure
    complex_data = {
        "metadata": {
            "version": "1.0",
            "created": datetime.now().isoformat()
        },
        "data": sample_data
    }
    
    with open(output_dir / 'complex_sample.json', 'w') as f:
        json.dump(complex_data, f, indent=2, default=str)
    
    logger.info(f"Sample files created in {output_dir}")